"""One-time import of the Commissioner's workbook (DFL_v2.xlsx) into plain CSV files the engine reads:

    python engine/import_workbook.py path/to/DFL_v2.xlsx

writes data/teams.csv (the 48 clubs: conference, division, seed or Exile, city, nickname, venue, location and three colours with hex codes where the workbook
has them) and data/schedule_2000.csv (every game of the first season with its date and kickoff, then the league-wide byes and the playoff dates). The engine
needs only the CSVs; openpyxl is needed only to run this importer.
"""
import csv
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _hex(cell):
    if not cell:
        return "", ""
    s = str(cell)
    if "\n#" in s:
        name, hx = s.split("\n#")
        return name.strip(), "#" + hx.strip()
    return s.strip(), ""


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    teams = list(wb["Teams"].iter_rows(values_only=True))
    head = [str(h) for h in teams[0]]
    ix = {h: i for i, h in enumerate(head)}
    colours = {}
    for r in list(wb["Changes"].iter_rows(values_only=True))[1:]:
        if r[0] is not None:
            colours[int(r[0])] = [_hex(r[8]), _hex(r[9]), _hex(r[10])]
    for r in list(wb["Expansion Colors"].iter_rows(values_only=True))[1:]:
        if r[0] is not None:
            colours[int(r[0])] = [_hex(r[9]), _hex(r[10]), _hex(r[11])]
    out = os.path.join(ROOT, "data", "teams.csv")
    n = 0
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["club_id", "conference", "division", "seed", "abbr", "city", "nickname", "venue", "location",
                    "primary", "primary_hex", "secondary", "secondary_hex", "tertiary", "tertiary_hex"])
        for r in teams[1:]:
            if r[ix["team_id"]] is None:
                continue
            cid = int(r[ix["team_id"]])
            c = colours.get(cid, [(r[ix["primary_color"]], ""), (r[ix["secondary_color"]], ""), ("", "")])
            seed = "Exile" if str(r[ix["schedule_creation_seeding"]]) == "Exile" else int(r[ix["current_division_rank"]])
            w.writerow([cid, r[ix["conference"]], r[ix["division"]], seed, r[ix["abreviation"]], r[ix["team_city"]], r[ix["team_name"]], r[ix["venue"]],
                        r[ix["location"]], c[0][0], c[0][1], c[1][0], c[1][1], c[2][0], c[2][1]])
            n += 1
    out2 = os.path.join(ROOT, "data", "schedule_2000.csv")
    m = 0
    with open(out2, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["season", "week", "game", "date", "kickoff", "away", "home", "note"])
        for r in list(wb["Schedule"].iter_rows(values_only=True))[1:]:
            if r[0] is not None and r[1] is not None:
                w.writerow([r[0], r[1], r[2], r[3].date().isoformat(), r[4].strftime("%H:%M"), r[5], r[6], ""])
                m += 1
            elif r[2] is not None and str(r[2]).startswith("League-Wide"):
                w.writerow(["", "", "", "", "", "", "", "league-wide bye"])
            elif r[3] is not None and r[5] is not None:
                w.writerow(["", "", "", r[3].date().isoformat(), "", "", "", str(r[5])])
            elif r[3] is not None and r[2] is None and r[4] is None:
                pass
    print(f"{n} clubs -> {out}; {m} games -> {out2}")


if __name__ == "__main__":
    main(sys.argv[1])
